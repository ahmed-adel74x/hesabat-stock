import os, tempfile, unittest
from datetime import datetime, timedelta
from helpers import tmp_db
from hesabat.db.connection import connect
from hesabat.db.migrations import init_db
from hesabat.logic import export, inventory, invoices as inv, parties, products, reports as rp, treasury as tr, warehouses
from hesabat.logic.errors import ValidationError


def day(n):
    return (datetime.now() - timedelta(days=n)).strftime("%Y-%m-%d")


class ReportBase(unittest.TestCase):
    def setUp(self):
        self.path = tmp_db(); self.c = connect(self.path); init_db(self.c, self.path)
        c = self.c
        self.w2 = warehouses.create(c, "فرع")
        self.sup = parties.create(c, "supplier", "مورد", opening_balance="100")
        self.cust = parties.create(c, "customer", "أحمد", opening_balance="50")
        self.cash = parties.default_customer_id(c)
        self.sugar = products.create(c, "سكر", "10", "18", "5", "10")     # 10 @ 10
        self.tea = products.create(c, "شاي", "5", "7", "3", "1")          # 1 @ 5 (تحت حد الطلب)
        inv.create_invoice(c, "purchase", self.sup, 1, [(self.sugar, "10", "20")], "0", "50", "", day(10))
        inv.create_invoice(c, "sale", self.cust, 1, [(self.sugar, "4", "20")], "0", "30", "", day(5))
        gone = inv.create_invoice(c, "sale", self.cash, 1, [(self.sugar, "1", "15")], "0", "15", "", day(2))
        inv.cancel_invoice(c, gone)                                        # فاتورة ملغية: لازم متظهرش
        inv.create_invoice(c, "sale", self.cash, 1, [(self.sugar, "2", "18")], "0", "36", "", day(1))
        inv.create_invoice(c, "sale_return", self.cust, 1, [(self.sugar, "1", "20")], "0", "0", "", day(1))
        inventory.transfer(c, 1, self.w2, [(self.sugar, "5")])
        tr.add_voucher(c, "payment", "50", category="كهرباء", entry_date=day(3))


class StockTests(ReportBase):
    def test_stock_all_warehouses(self):
        r = rp.stock_report(self.c)
        self.assertEqual([t for t, _k in r["columns"]], ["الصنف", "المخزن الرئيسي", "فرع", "الإجمالي", "حد التنبيه",
                                                         "الحالة", "القيمة بالتكلفة", "القيمة بسعر البيع"])
        self.assertEqual(r["rows"][0], ["سكر", 10000, 5000, 15000, 5000, "متوفر", 22500, 27000])
        self.assertEqual(r["rows"][1], ["شاي", 1000, 0, 1000, 3000, "وصل لحد الطلب", 500, 700])
        self.assertEqual([v for _l, _k, v in r["summary"]], [2, 1, 23000, 27700])

    def test_stock_one_warehouse_and_low(self):
        r = rp.stock_report(self.c, self.w2)
        self.assertEqual(len(r["columns"]), 6)
        self.assertEqual(r["rows"][0][:4], ["سكر", 5000, 5000, "وصل لحد الطلب"])      # الكمية = حد التنبيه
        with self.assertRaises(ValidationError):
            rp.stock_report(self.c, 999)
        low = rp.low_stock_report(self.c)
        self.assertEqual(low["rows"], [["شاي", 1000, 3000, 2000]])


class InvoiceReportTests(ReportBase):
    def test_sales(self):
        r = rp.sales_report(self.c)
        self.assertEqual([(x[0], x[2], x[3], x[4], x[5], x[6]) for x in r["rows"]], [
            (day(5), "فاتورة بيع", "أحمد", 8000, 3000, 5000),
            (day(1), "فاتورة بيع", "عميل نقدي", 3600, 3600, 0),
            (day(1), "مرتجع بيع", "أحمد", -2000, 0, -2000)])
        self.assertEqual([v for _l, _k, v in r["summary"]], [2, 11600, 2000, 9600, 6600, 3000])
        part = rp.sales_report(self.c, day(2), day(0))                         # الفاتورة الملغية مش موجودة برضه
        self.assertEqual(len(part["rows"]), 2)
        self.assertEqual(rp.sales_report(self.c, day(30), day(20))["rows"], [])
        with self.assertRaises(ValidationError):
            rp.sales_report(self.c, day(1), day(5))
        with self.assertRaises(ValidationError):
            rp.sales_report(self.c, "مش تاريخ")

    def test_purchases(self):
        r = rp.purchases_report(self.c)
        self.assertEqual(r["rows"], [[day(10), 1, "فاتورة شراء", "مورد", 20000, 5000, 15000]])
        self.assertEqual([v for _l, _k, v in r["summary"]], [1, 20000, 0, 20000, 5000, 15000])

    def test_profit(self):
        r = rp.profit_report(self.c)
        self.assertEqual(r["rows"], [["سكر", 5000, 9600, 7500, 2100, "21.9%"]])   # 4+2-1 مرتجع، والملغية اتعادلت
        self.assertEqual([v for _l, _k, v in r["summary"]], [9600, 7500, 2100, 5000, -2900])
        part = rp.profit_report(self.c, day(2), day(0))
        self.assertEqual(part["rows"][0][2:5], [1600, 1500, 100])
        self.assertEqual(part["summary"][3][2], 0)                               # مصروف الكهرباء برا الفترة


class StatementTests(ReportBase):
    def test_customer_statement(self):
        r = rp.statement_report(self.c, "customer", self.cust)
        self.assertEqual([x[4] for x in r["rows"]], [5000, 13000, 10000, 8000])
        self.assertEqual([v for _l, _k, v in r["summary"]], [5000, 3000, 8000])
        self.assertEqual({p["id"]: p["balance_c"] for p in parties.list_parties(self.c, "customer")}[self.cust], 8000)
        part = rp.statement_report(self.c, "customer", self.cust, day(2), day(0))
        self.assertEqual([v for _l, _k, v in part["summary"]], [10000, -2000, 8000])   # رصيد أول المدة بعد حركة قبل الفترة
        self.assertEqual(len(part["rows"]), 2)
        cash = rp.statement_report(self.c, "customer", self.cash)
        self.assertEqual(cash["summary"][2][2], 0)                                # العميل النقدي رصيده صفر

    def test_supplier_statement(self):
        r = rp.statement_report(self.c, "supplier", self.sup)
        self.assertEqual([x[4] for x in r["rows"]], [10000, 30000, 25000])
        with self.assertRaises(ValidationError):
            rp.statement_report(self.c, "supplier", 9999)


class MovementAndTreasuryTests(ReportBase):
    def test_product_movement(self):
        r = rp.product_movement_report(self.c, self.sugar)
        self.assertEqual(r["summary"][0][2], 0)
        self.assertEqual(r["summary"][3][2], inventory.balance(self.c, self.sugar))
        self.assertEqual(r["rows"][-1][6], 15000)
        self.assertEqual(r["summary"][1][2] - r["summary"][2][2], 15000)
        main = rp.product_movement_report(self.c, self.sugar, 1)
        self.assertEqual(main["summary"][3][2], inventory.balance(self.c, self.sugar, 1))
        part = rp.product_movement_report(self.c, self.sugar, None, day(1), day(0))
        self.assertEqual(part["summary"][0][2], 6000)                            # الرصيد قبل الفترة
        self.assertEqual(part["summary"][3][2], 15000)
        kinds = {x[2] for x in r["rows"]}
        self.assertTrue({"شراء", "بيع", "مرتجع بيع", "رصيد افتتاحي"} <= kinds)

    def test_treasury(self):
        r = rp.treasury_report(self.c)
        self.assertEqual(r["summary"][3][2], tr.balance(self.c))
        self.assertEqual([x[5] for x in r["rows"]], [0, -5000, -2000, -7000, -3400])
        self.assertEqual(len(r["rows"]), 5)                                      # الدفعة الملغية مش ظاهرة

    def test_registry(self):
        for key, title, params in rp.REPORTS:
            p = {"customer_id": self.cust, "supplier_id": self.sup, "product_id": self.sugar}
            rep = rp.run_report(self.path, key, p)                                # بيفتح اتصاله الخاص
            self.assertEqual(len(rep["columns"]), len(rep["rows"][0]) if rep["rows"] else len(rep["columns"]))
        with self.assertRaises(ValidationError):
            rp.run_report(self.path, "customer_statement", {})
        with self.assertRaises(ValidationError):
            rp.run_report(self.path, "product_movement", {})
        with self.assertRaises(ValueError):
            rp.run_report(self.path, "nope", {})

    def test_format_cell(self):
        self.assertEqual(rp.format_cell("money", -150000), "-1,500.00")
        self.assertEqual(rp.format_cell("qty", 2500), "2.5")
        self.assertEqual(rp.format_cell("money", None), "")
        self.assertEqual(rp.format_cell("text", "x"), "x")


class ExportTests(ReportBase):
    def test_excel_roundtrip(self):
        from openpyxl import load_workbook
        r = rp.sales_report(self.c)
        path = os.path.join(tempfile.mkdtemp(), "s.xlsx")
        export.to_excel(r, path)
        ws = load_workbook(path).active
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws["A1"].value, "تقرير المبيعات")
        self.assertEqual([c.value for c in ws[4]][:4], ["التاريخ", "الرقم", "النوع", "العميل"])
        self.assertEqual(ws["E5"].value, 80.0)                                   # أرقام حقيقية مش نصوص
        self.assertEqual(ws["E5"].number_format, "#,##0.00")
        self.assertEqual(ws["E7"].value, -20.0)
        labels = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(9, 16)}
        self.assertEqual(labels["الصافي"], 96.0)

    def test_excel_sheet_title_and_formula_safety(self):
        from openpyxl import load_workbook
        products.create(self.c, "=1+1", "1", "2")
        path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
        export.to_excel(rp.stock_report(self.c), path)
        cell = next(c for c in load_workbook(path).active["A"] if c.value == "=1+1")
        self.assertEqual(cell.data_type, "s")                                    # نص مش معادلة
        path2 = os.path.join(tempfile.mkdtemp(), "c.xlsx")
        export.to_excel(rp.statement_report(self.c, "customer", self.cust), path2)   # العنوان فيه ":" (ممنوعة في اسم الورقة)
        self.assertEqual(load_workbook(path2).active.title.count(":"), 0)

    def test_default_filename(self):
        n = export.default_filename({"title": "كشف حساب العميل: أحمد / 1"}, "xlsx")
        self.assertTrue(n.endswith(".xlsx"))
        for ch in '\\/:*?"<>| ':
            self.assertNotIn(ch, n)

    def test_html(self):
        products.create(self.c, "<script>x</script>", "1", "2")
        h = export.report_html(rp.stock_report(self.c))
        self.assertIn("تقرير المخزون", h)
        self.assertIn("&lt;script&gt;", h)
        self.assertNotIn("<script>", h)
        self.assertIn('dir="rtl"', h)


if __name__ == "__main__":
    unittest.main(verbosity=2)
