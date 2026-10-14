"""تصدير التقارير: Excel (openpyxl) ونص HTML (بتطبعه الواجهة PDF أو على الطابعة)."""
import html
import re
from decimal import Decimal

from .errors import ValidationError
from .reports import format_cell


def _value(kind, v):
    """قيمة الخلية في Excel: أرقام حقيقية (مش نص) عشان تتجمع وتتفلتر."""
    if v is None or v == "":
        return None
    if kind == "money":
        return float(Decimal(v) / 100)
    if kind == "qty":
        return float(Decimal(v) / 1000)
    if kind == "int":
        return int(v)
    return str(v)


def to_excel(report, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ValidationError("مكتبة openpyxl مش متثبتة. نفّذ في الترمينال: pip install openpyxl")
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\\/*?:\[\]]", " ", report["title"])[:31]) or "تقرير"
    ws.sheet_view.rightToLeft = True
    cols = report["columns"]

    def put(row, col, val, kind="text", bold=False):
        c = ws.cell(row=row, column=col)
        c.value = _value(kind, val)
        if isinstance(c.value, str) and c.value[:1] == "=":
            c.data_type = "s"          # نص مش معادلة (حماية من حقن المعادلات)
        if kind == "money":
            c.number_format = "#,##0.00"
        elif kind == "qty":
            c.number_format = "#,##0.###"
        if bold:
            c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
        return c

    ws.cell(row=1, column=1, value=report["title"]).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=report["subtitle"])
    r = 4
    for j, (title, _k) in enumerate(cols, 1):
        c = put(r, j, title, bold=True)
        c.fill = PatternFill("solid", fgColor="E2E8F0")
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    for row in report["rows"]:
        r += 1
        for j, ((_t, kind), v) in enumerate(zip(cols, row), 1):
            put(r, j, v, kind)
    r += 2
    for label, kind, v in report["summary"]:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        put(r, 2, v, kind, bold=True)
        r += 1
    for j in range(1, len(cols) + 1):
        width = max([len(str(cols[j - 1][0]))] + [len(format_cell(cols[j - 1][1], row[j - 1])) for row in report["rows"]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 4, 10), 50)
    wb.save(path)


def report_html(report):
    esc = html.escape
    head = "".join(f"<th>{esc(t)}</th>" for t, _k in report["columns"])
    body = "".join("<tr>" + "".join(f"<td align='center'>{esc(format_cell(k, v))}</td>"
                                    for (_t, k), v in zip(report["columns"], row)) + "</tr>" for row in report["rows"])
    summ = "".join(f"<tr><td><b>{esc(lbl)}</b></td><td align='center'>{esc(format_cell(k, v))}</td></tr>"
                   for lbl, k, v in report["summary"])
    return f"""<html dir="rtl"><body dir="rtl" style="font-size:10pt">
<h2 align="center">{esc(report['title'])}</h2><p align="center">{esc(report['subtitle'])}</p>
<table dir="rtl" width="100%" border="1" cellspacing="0" cellpadding="4">
<tr bgcolor="#eeeeee">{head}</tr>{body}</table><br>
<table dir="rtl" border="0" cellpadding="3">{summ}</table></body></html>"""


def default_filename(report, ext):
    """اسم ملف مقترح من عنوان التقرير والتاريخ (من غير رموز ممنوعة في أسماء الملفات)."""
    from datetime import datetime
    name = re.sub(r'[\\/*?:"<>|\s]+', "_", report["title"]).strip("_")
    return f"{name}_{datetime.now():%Y-%m-%d}.{ext}"
